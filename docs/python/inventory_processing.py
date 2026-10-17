
import sys
import math
from pathlib import Path
from datetime import datetime
from collections import defaultdict

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
import argparse

# ==========================================
# 1. CAU HINH
# ==========================================

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(
        encoding="utf-8",
        errors="replace"
    )

parser = argparse.ArgumentParser(
    description="Xử lý dữ liệu tồn kho"
)

parser.add_argument(
    "input",
    help="Đường dẫn file Excel đầu vào"
)

parser.add_argument(
    "output",
    help="Đường dẫn file Excel kết quả"
)

parser.add_argument(
    "--sheet-ma",
    default="mã rượu",
    help="Tên sheet cần gán mã"
)

parser.add_argument(
    "--sheet-ton",
    default="tồn rượu",
    help="Tên sheet tồn kho"
)

args = parser.parse_args()

INPUT_FILE = Path(args.input)
OUTPUT_FILE = Path(args.output)

SHEET_MA = args.sheet_ma
SHEET_TON = args.sheet_ton
SHEET_LOG = "LOG_XU_LY"
SHEET_TONG_HOP = "TONG_HOP_TON_KHO"

# Sheet ma ruou
COL_TEN_MA = 1       # A - Ten hang
COL_MA_GAN = 2       # B - Ma VTHH
COL_QUY_DOI = 5      # E - Quy doi

# Sheet ton ruou
COL_MA_TON = 2       # B - Ma hang
COL_TEN_TON = 3      # C - Ten hang
COL_SO_LUONG = 4     # D - So luong ton kho

START_ROW_MA = 2
START_ROW_TON = 6

# Cho phep trung ma neu ton kho cua ma do van du
ALLOW_DUPLICATE_CODE = True

# Xoa ma cu va gan lai
RESET_OLD_CODES = True

# Canh bao neu so luong du sau gan > 20% quy doi
WARNING_PERCENT = 20

# Dung de so sanh so thuc
EPSILON = 1e-9


# ==========================================
# 2. BIEN THONG KE
# ==========================================

logs = []
ton_kho = {}

so_thanh_cong = 0
so_canh_bao = 0
so_loi = 0
so_bo_qua = 0

# Luu danh sach cac dong duoc gan theo tung ma
lich_su_gan = defaultdict(list)


# ==========================================
# 3. HAM CHUYEN DOI
# ==========================================

def to_number(value):
    if value is None:
        return None

    try:
        if isinstance(value, bool):
            return None

        if isinstance(value, (int, float)):
            number = float(value)
        else:
            value = str(value).strip()

            if not value:
                return None

            value = value.replace(",", "")
            number = float(value)

        if not math.isfinite(number):
            return None

        return number

    except (ValueError, TypeError):
        return None


def normalize_code(value):
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


def format_number(value):
    if value is None:
        return "-"

    if isinstance(value, (int, float)):
        return f"{value:,.2f}".rstrip("0").rstrip(".")

    return str(value)


# ==========================================
# 4. HAM GHI LOG CHI TIET
# ==========================================

def add_log(
    status,
    sheet,
    row,
    ten_hang="",
    quy_doi=None,
    ma_hang="",
    ton_ban_dau=None,
    ton_truoc=None,
    ton_sau=None,
    ly_do=""
):
    global so_canh_bao, so_loi, so_bo_qua

    if status == "ERROR":
        so_loi += 1
    elif status == "WARNING":
        so_canh_bao += 1
    elif status == "SKIP":
        so_bo_qua += 1

    log = {
        "status": status,
        "sheet": sheet,
        "row": row,
        "ten_hang": ten_hang,
        "quy_doi": quy_doi,
        "ma_hang": ma_hang,
        "ton_ban_dau": ton_ban_dau,
        "ton_truoc": ton_truoc,
        "ton_sau": ton_sau,
        "ly_do": ly_do
    }

    logs.append(log)

    print(
        f"[{status}] "
        f"Sheet: {sheet} | "
        f"Dong: {row} | "
        f"Ten: {ten_hang} | "
        f"Quy doi: {format_number(quy_doi)} | "
        f"Ma: {ma_hang or '-'} | "
        f"Ton ban dau: {format_number(ton_ban_dau)} | "
        f"Ton truoc: {format_number(ton_truoc)} | "
        f"Ton sau: {format_number(ton_sau)} | "
        f"Chi tiet: {ly_do}"
    )


# ==========================================
# 5. DOC FILE EXCEL
# ==========================================

if not INPUT_FILE.exists():
    raise FileNotFoundError(
        f"Khong tim thay file: {INPUT_FILE}"
    )

# Workbook ghi ket qua, giu format goc
wb = load_workbook(INPUT_FILE)

# Workbook doc gia tri da tinh cua cong thuc
wb_data = load_workbook(
    INPUT_FILE,
    data_only=True
)

for sheet_name in [SHEET_MA, SHEET_TON]:
    if sheet_name not in wb.sheetnames:
        raise ValueError(
            f"Khong tim thay sheet: {sheet_name}"
        )

ws_ma = wb[SHEET_MA]
ws_ma_data = wb_data[SHEET_MA]
ws_ton = wb_data[SHEET_TON]

print("=" * 80)
print("BAT DAU XU LY TON KHO RUOU")
print("=" * 80)

print(f"File: {INPUT_FILE.name}")
print(f"Thoi gian: {datetime.now()}")
print(f"Cho phep trung ma: {ALLOW_DUPLICATE_CODE}")


# ==========================================
# 6. DOC TON KHO THEO TUNG MA
# ==========================================

print("\n" + "=" * 80)
print("1. DOC TON KHO THEO TUNG MA")
print("=" * 80)

for row in range(
    START_ROW_TON,
    ws_ton.max_row + 1
):
    ma_hang = normalize_code(
        ws_ton.cell(row, COL_MA_TON).value
    )

    ten_hang = ws_ton.cell(
        row, COL_TEN_TON
    ).value

    so_luong_goc = ws_ton.cell(
        row, COL_SO_LUONG
    ).value

    so_luong = to_number(so_luong_goc)

    # Bo qua dong hoan toan trong
    if (
        not ma_hang
        and ten_hang is None
        and so_luong_goc is None
    ):
        continue

    if not ma_hang:
        add_log(
            "ERROR",
            SHEET_TON,
            row,
            ten_hang,
            ly_do="Ma hang cot B bi trong"
        )
        continue

    if so_luong is None:
        add_log(
            "ERROR",
            SHEET_TON,
            row,
            ten_hang,
            ma_hang=ma_hang,
            ly_do=(
                f"So luong cot D khong hop le: "
                f"{so_luong_goc}"
            )
        )
        continue

    if so_luong < 0:
        add_log(
            "ERROR",
            SHEET_TON,
            row,
            ten_hang,
            ma_hang=ma_hang,
            ly_do="So luong ton kho khong duoc am"
        )
        continue

    if ma_hang in ton_kho:
        add_log(
            "ERROR",
            SHEET_TON,
            row,
            ten_hang,
            ma_hang=ma_hang,
            ly_do=(
                f"Ma ton kho bi trung voi dong "
                f"{ton_kho[ma_hang]['row']}. "
                "Bo qua dong trung, khong cong don ton."
            )
        )
        continue

    # Moi ma co ton kho RIENG
    ton_kho[ma_hang] = {
        "row": row,
        "ma_hang": ma_hang,
        "ten_hang": ten_hang,
        "ton_ban_dau": so_luong,
        "da_gan": 0.0,
        "con_lai": so_luong,
        "so_lan_gan": 0
    }

print(f"\nTong ma ton kho hop le: {len(ton_kho)}")

if not ton_kho:
    raise ValueError(
        "Khong co ma ton kho hop le. "
        "Kiem tra cot B va D cua sheet ton ruou."
    )


# ==========================================
# 7. DOC DANH SACH CAN GAN MA
# ==========================================

print("\n" + "=" * 80)
print("2. DOC DANH SACH CAN GAN MA")
print("=" * 80)

danh_sach_can_gan = []

for row in range(
    START_ROW_MA,
    ws_ma.max_row + 1
):
    ten_hang_goc = ws_ma.cell(
        row, COL_TEN_MA
    ).value

    gia_tri_goc = ws_ma_data.cell(
        row, COL_QUY_DOI
    ).value

    # Bo qua dong trong
    if (
        ten_hang_goc is None
        and gia_tri_goc is None
    ):
        continue

    ten_hang = (
        str(ten_hang_goc).strip()
        if ten_hang_goc is not None
        else ""
    )

    # Xoa ma cu de gan lai tu dau
    if RESET_OLD_CODES:
        ws_ma.cell(row, COL_MA_GAN).value = None

    if not ten_hang:
        add_log(
            "ERROR",
            SHEET_MA,
            row,
            quy_doi=gia_tri_goc,
            ly_do="Ten hang cot A bi trong"
        )
        continue

    quy_doi = to_number(gia_tri_goc)

    if quy_doi is None:
        add_log(
            "ERROR",
            SHEET_MA,
            row,
            ten_hang,
            quy_doi=gia_tri_goc,
            ly_do=(
                "Gia tri quy doi cot E trong "
                "hoac khong phai so hop le"
            )
        )
        continue

    if quy_doi <= 0:
        add_log(
            "ERROR",
            SHEET_MA,
            row,
            ten_hang,
            quy_doi,
            ly_do="Gia tri quy doi phai lon hon 0"
        )
        continue

    danh_sach_can_gan.append({
        "row": row,
        "ten_hang": ten_hang,
        "quy_doi": quy_doi
    })

print(
    f"\nTong dong can gan: "
    f"{len(danh_sach_can_gan)}"
)


# ==========================================
# 8. GAN MA THEO TON KHO RIENG
# ==========================================

print("\n" + "=" * 80)
print("3. GAN MA THEO TON KHO CON LAI")
print("=" * 80)

for item_ma in danh_sach_can_gan:

    row = item_ma["row"]
    ten_hang = item_ma["ten_hang"]
    quy_doi = item_ma["quy_doi"]

    # Chi tim nhung ma con DU ton kho
    danh_sach_kha_dung = [
        item
        for item in ton_kho.values()
        if (
            item["con_lai"] + EPSILON >= quy_doi
            and (
                ALLOW_DUPLICATE_CODE
                or item["so_lan_gan"] == 0
            )
        )
    ]

    if not danh_sach_kha_dung:

        ton_lon_nhat = max(
            (
                item["con_lai"]
                for item in ton_kho.values()
            ),
            default=0
        )

        add_log(
            "ERROR",
            SHEET_MA,
            row,
            ten_hang,
            quy_doi,
            ly_do=(
                "Khong co ma nao du ton kho rieng "
                f"cho quy doi {format_number(quy_doi)}. "
                f"Ton rieng lon nhat con lai: "
                f"{format_number(ton_lon_nhat)}. "
                "Khong duoc cong don ton cac ma."
            )
        )
        continue

    # Chon ma co ton CON LAI vua du nhat
    # Neu bang nhau, uu tien ma da duoc su dung
    # de tranh phan tan qua nhieu ma
    gan_nhat = min(
        danh_sach_kha_dung,
        key=lambda item: (
            item["con_lai"] - quy_doi,
            -item["so_lan_gan"],
            item["ma_hang"]
        )
    )

    ma_hang = gan_nhat["ma_hang"]
    ton_ban_dau = gan_nhat["ton_ban_dau"]
    ton_truoc = gan_nhat["con_lai"]

    # Kiem tra an toan truoc khi tru ton
    if ton_truoc + EPSILON < quy_doi:
        add_log(
            "ERROR",
            SHEET_MA,
            row,
            ten_hang,
            quy_doi,
            ma_hang,
            ton_ban_dau,
            ton_truoc,
            ton_truoc,
            "Ton kho khong du, huy gan ma"
        )
        continue

    # TRU TON CUA CHINH MA HANG NAY
    gan_nhat["da_gan"] += quy_doi

    gan_nhat["con_lai"] = max(
        0.0,
        ton_ban_dau - gan_nhat["da_gan"]
    )

    gan_nhat["so_lan_gan"] += 1

    ton_sau = gan_nhat["con_lai"]

    # Gan ma vao cot B
    ws_ma.cell(
        row, COL_MA_GAN
    ).value = ma_hang

    # Luu lich su theo tung ma
    lich_su_gan[ma_hang].append({
        "row": row,
        "ten_hang": ten_hang,
        "quy_doi": quy_doi,
        "ton_truoc": ton_truoc,
        "ton_sau": ton_sau
    })

    so_thanh_cong += 1

    # Ty le ton du so voi quy doi
    ty_le_du = (
        ton_sau / quy_doi * 100
    )

    if ty_le_du > WARNING_PERCENT:
        status = "WARNING"
        ly_do = (
            f"Gan thanh cong. Ton du sau gan "
            f"lon hon {WARNING_PERCENT}% quy doi. "
            f"Lan gan thu {gan_nhat['so_lan_gan']} "
            f"cua ma {ma_hang}."
        )
    else:
        status = "SUCCESS"
        ly_do = (
            f"Gan thanh cong. "
            f"Lan gan thu {gan_nhat['so_lan_gan']} "
            f"cua ma {ma_hang}."
        )

    add_log(
        status,
        SHEET_MA,
        row,
        ten_hang,
        quy_doi,
        ma_hang,
        ton_ban_dau,
        ton_truoc,
        ton_sau,
        ly_do
    )


# ==========================================
# 9. KIEM TRA TON KHO TUNG MA
# ==========================================

print("\n" + "=" * 80)
print("4. KIEM TRA TON KHO TUNG MA")
print("=" * 80)

so_ma_vuot_ton = 0

for ma_hang, item in ton_kho.items():

    if item["so_lan_gan"] == 0:
        continue

    ton_ban_dau = item["ton_ban_dau"]
    da_gan = item["da_gan"]
    con_lai = item["con_lai"]

    if da_gan > ton_ban_dau + EPSILON:
        so_ma_vuot_ton += 1
        status = "ERROR"
    else:
        status = "OK"

    print(
        f"[{status}] "
        f"Ma: {ma_hang} | "
        f"Ton ban dau: {format_number(ton_ban_dau)} | "
        f"So lan gan: {item['so_lan_gan']} | "
        f"Tong quy doi da gan: {format_number(da_gan)} | "
        f"Ton con lai: {format_number(con_lai)}"
    )

if so_ma_vuot_ton:
    raise ValueError(
        f"Co {so_ma_vuot_ton} ma vuot ton kho. "
        "Dung xu ly, khong luu file."
    )


# ==========================================
# 10. HAM DINH DANG SHEET BAO CAO
# ==========================================

def format_report_sheet(ws, widths):

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78"
    )

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    ws.row_dimensions[1].height = 32

    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


# ==========================================
# 11. TAO SHEET LOG_XU_LY
# ==========================================

print("\n" + "=" * 80)
print("5. TAO SHEET LOG_XU_LY")
print("=" * 80)

if SHEET_LOG in wb.sheetnames:
    del wb[SHEET_LOG]

ws_log = wb.create_sheet(SHEET_LOG)

ws_log.append([
    "Trạng thái",
    "Sheet",
    "Dòng Excel",
    "Tên hàng",
    "Quy đổi",
    "Mã hàng",
    "Tồn ban đầu",
    "Tồn trước khi gán",
    "Tồn sau khi gán",
    "Chi tiết"
])

for item in logs:
    ws_log.append([
        item["status"],
        item["sheet"],
        item["row"],
        item["ten_hang"],
        item["quy_doi"],
        item["ma_hang"],
        item["ton_ban_dau"],
        item["ton_truoc"],
        item["ton_sau"],
        item["ly_do"]
    ])

colors = {
    "SUCCESS": "E2F0D9",
    "WARNING": "FFF2CC",
    "ERROR": "FCE4D6",
    "SKIP": "E7E6E6"
}

for row in range(2, ws_log.max_row + 1):
    status = ws_log.cell(row, 1).value

    ws_log.cell(row, 1).fill = PatternFill(
        "solid",
        fgColor=colors.get(status, "FFFFFF")
    )

    for col in [5, 7, 8, 9]:
        ws_log.cell(row, col).number_format = "#,##0.00"

format_report_sheet(
    ws_log,
    {
        "A": 15,
        "B": 18,
        "C": 14,
        "D": 45,
        "E": 18,
        "F": 18,
        "G": 20,
        "H": 22,
        "I": 22,
        "J": 75
    }
)


# ==========================================
# 12. TAO SHEET TONG_HOP_TON_KHO
# ==========================================

print("\n" + "=" * 80)
print("6. TAO SHEET TONG_HOP_TON_KHO")
print("=" * 80)

if SHEET_TONG_HOP in wb.sheetnames:
    del wb[SHEET_TONG_HOP]

ws_tong_hop = wb.create_sheet(SHEET_TONG_HOP)

ws_tong_hop.append([
    "Mã hàng",
    "Tên hàng",
    "Dòng tồn kho",
    "Tồn ban đầu",
    "Số lần gán",
    "Tổng quy đổi đã gán",
    "Tồn còn lại",
    "Các dòng đã gán",
    "Trạng thái"
])

for ma_hang, item in ton_kho.items():

    rows_gan = [
        str(x["row"])
        for x in lich_su_gan[ma_hang]
    ]

    if item["da_gan"] > item["ton_ban_dau"] + EPSILON:
        status = "VUOT TON"
    elif item["so_lan_gan"] == 0:
        status = "CHUA SU DUNG"
    elif item["con_lai"] <= EPSILON:
        status = "HET TON"
    else:
        status = "CON TON"

    ws_tong_hop.append([
        ma_hang,
        item["ten_hang"],
        item["row"],
        item["ton_ban_dau"],
        item["so_lan_gan"],
        item["da_gan"],
        item["con_lai"],
        ", ".join(rows_gan),
        status
    ])

for row in range(2, ws_tong_hop.max_row + 1):

    for col in [4, 6, 7]:
        ws_tong_hop.cell(
            row, col
        ).number_format = "#,##0.00"

    status = ws_tong_hop.cell(row, 9).value

    status_colors = {
        "CON TON": "E2F0D9",
        "HET TON": "FFF2CC",
        "CHUA SU DUNG": "E7E6E6",
        "VUOT TON": "FCE4D6"
    }

    ws_tong_hop.cell(row, 9).fill = PatternFill(
        "solid",
        fgColor=status_colors.get(status, "FFFFFF")
    )

format_report_sheet(
    ws_tong_hop,
    {
        "A": 18,
        "B": 50,
        "C": 18,
        "D": 20,
        "E": 18,
        "F": 25,
        "G": 20,
        "H": 45,
        "I": 20
    }
)


# ==========================================
# 13. LUU FILE EXCEL
# ==========================================

print("\n" + "=" * 80)
print("7. LUU FILE KET QUA")
print("=" * 80)

try:
    wb.save(OUTPUT_FILE)

except PermissionError:
    raise PermissionError(
        f"Khong the luu file {OUTPUT_FILE}. "
        "Hay dong file Excel dang mo."
    )


# ==========================================
# 14. TONG KET
# ==========================================

so_ma_da_dung = sum(
    1 for item in ton_kho.values()
    if item["so_lan_gan"] > 0
)

so_ma_gan_nhieu_lan = sum(
    1 for item in ton_kho.values()
    if item["so_lan_gan"] > 1
)

so_dong_chua_gan = sum(
    1
    for item in logs
    if (
        item["status"] == "ERROR"
        and item["sheet"] == SHEET_MA
    )
)

print("\n" + "=" * 80)
print("KET QUA CUOI CUNG")
print("=" * 80)

print(f"Tong ma ton kho: {len(ton_kho)}")
print(f"Tong dong hop le can gan: {len(danh_sach_can_gan)}")
print(f"So dong gan thanh cong: {so_thanh_cong}")
print(f"So dong loi chua gan: {so_dong_chua_gan}")

print(f"So log canh bao: {so_canh_bao}")
print(f"So log loi: {so_loi}")

print(f"So ma da su dung: {so_ma_da_dung}")
print(f"So ma duoc gan nhieu lan: {so_ma_gan_nhieu_lan}")
print(f"So ma vuot ton: {so_ma_vuot_ton}")

print("\nCHI TIET MA DUOC GAN NHIEU LAN:")

for ma_hang, item in ton_kho.items():

    if item["so_lan_gan"] <= 1:
        continue

    print(
        f"Ma: {ma_hang} | "
        f"So lan: {item['so_lan_gan']} | "
        f"Ton ban dau: {format_number(item['ton_ban_dau'])} | "
        f"Da gan: {format_number(item['da_gan'])} | "
        f"Con lai: {format_number(item['con_lai'])}"
    )

print(f"\nFile ket qua: {OUTPUT_FILE}")
print("\nHOAN THANH!")
